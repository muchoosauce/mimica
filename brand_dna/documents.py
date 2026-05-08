"""Parse PDF / DOCX brand-guideline documents into text + images."""
from __future__ import annotations

import shutil
from pathlib import Path

from .types import CollectedImage, ParsedDocument


def parse_document(path: Path, image_dir: Path) -> ParsedDocument:
    """Dispatch on extension. Returns ParsedDocument with .error set on failure."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _parse_pdf(path, image_dir)
    if ext in (".docx",):
        return _parse_docx(path, image_dir)
    if ext in (".txt", ".md"):
        try:
            return ParsedDocument(path=path, text=path.read_text(encoding="utf-8", errors="ignore"))
        except Exception as e:
            return ParsedDocument(path=path, error=f"read failed: {e}")
    return ParsedDocument(path=path, error=f"unsupported extension {ext!r}")


def _parse_pdf(path: Path, image_dir: Path) -> ParsedDocument:
    try:
        from pypdf import PdfReader
    except ImportError as e:
        return ParsedDocument(path=path, error=f"pypdf not installed: {e}")

    images: list[CollectedImage] = []
    text_parts: list[str] = []
    try:
        reader = PdfReader(str(path))
        for page_idx, page in enumerate(reader.pages):
            try:
                text_parts.append(page.extract_text() or "")
            except Exception:
                pass
            try:
                for img_idx, img in enumerate(getattr(page, "images", []) or []):
                    out = image_dir / f"{path.stem}_p{page_idx:02d}_{img_idx:02d}.png"
                    with open(out, "wb") as f:
                        f.write(img.data)
                    images.append(CollectedImage(
                        path=out,
                        origin=f"document:{path.name}",
                        note=f"page {page_idx + 1}",
                    ))
            except Exception:
                pass
    except Exception as e:
        return ParsedDocument(path=path, error=f"pdf parse failed: {e}")

    return ParsedDocument(path=path, text="\n\n".join(text_parts).strip(), images=images)


def _parse_docx(path: Path, image_dir: Path) -> ParsedDocument:
    try:
        import docx  # python-docx
    except ImportError as e:
        return ParsedDocument(path=path, error=f"python-docx not installed: {e}")

    images: list[CollectedImage] = []
    try:
        document = docx.Document(str(path))
        text_parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_parts.append(cell.text.strip())
        # Embedded images: walk relationships for image parts
        rels = document.part.rels
        for rel_id, rel in rels.items():
            target = rel.target_part if hasattr(rel, "target_part") else None
            if target is None:
                continue
            content_type = getattr(target, "content_type", "") or ""
            if not content_type.startswith("image/"):
                continue
            ext = content_type.split("/")[-1] or "png"
            if ext == "jpeg":
                ext = "jpg"
            out = image_dir / f"{path.stem}_{rel_id}.{ext}"
            try:
                with open(out, "wb") as f:
                    f.write(target.blob)
                images.append(CollectedImage(
                    path=out,
                    origin=f"document:{path.name}",
                    note="embedded",
                ))
            except Exception:
                continue
    except Exception as e:
        return ParsedDocument(path=path, error=f"docx parse failed: {e}")

    return ParsedDocument(path=path, text="\n".join(text_parts).strip(), images=images)


def stage_creative(path: Path, image_dir: Path) -> CollectedImage:
    """Copy a creative file into the run's image dir as a CollectedImage."""
    dst = image_dir / path.name
    if path.resolve() != dst.resolve():
        shutil.copy2(path, dst)
    return CollectedImage(path=dst, origin="creative", note="user-provided")
